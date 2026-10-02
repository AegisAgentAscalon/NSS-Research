from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json,re
R=Path('/mnt/data/NSS_LB03_LEARNABILITY_v0.1');sections=json.loads((R/'docs/REPORT_SECTIONS.json').read_text())
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10);style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.08
for n,size,col in [('Title',28,'15354A'),('Heading 1',19,'15354A'),('Heading 2',12,'1E6971'),('Heading 3',10.5,'15354A')]:
 st=doc.styles[n];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(col);st.paragraph_format.space_after=Pt(10);st.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0];header.text='NEURAL SPACE SYSTEMS  /  RESEARCH RECORD';header.style='Caption';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('566875')
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('LB03-LEARNABILITY  •  v0.1   |   ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)

def rich(p,text):
 text=re.sub(r'`([^`]+)`',r'\1',text)
 for i,part in enumerate(re.split(r'(\*\*.*?\*\*)',text)):
  if part.startswith('**'):r=p.add_run(part[2:-2]);r.bold=True
  else:r=p.add_run(part)
 return p

def bookmark(p,name):
 bid=str(200+int(name[3:]));a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),bid);a.set(qn('w:name'),name);b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),bid);p._p.insert(0,a);p._p.append(b)
def link(p,text,anchor):
 a=OxmlElement('w:hyperlink');a.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1E6971');props.append(c);r.append(props);t=OxmlElement('w:t');t.text=text;r.append(t);a.append(r);p._p.append(a)
def table(lines):
 rows=[[x.strip()for x in line.strip().strip('|').split('|')]for line in lines if not re.match(r'^\s*\|?[- :|]+\|?\s*$',line)]
 if not rows:return
 cols=len(rows[0]);tb=doc.add_table(rows=1,cols=cols);tb.autofit=False
 widths={2:[2.15,5.05],3:[2.0,2.45,2.75],4:[2.45,1.58,1.58,1.59],5:[2.15,1.40,1.45,1.1,1.1]}.get(cols,[7.2/cols]*cols)
 for c,w in zip(tb.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  rr=tb.rows[0]if i==0 else tb.add_row()
  for j,text in enumerate(row):
   cell=rr.cells[j];cell.width=Inches(widths[j]);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.0;rich(p,text)
   for r in p.runs:r.font.size=Pt(9)
   tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'15354A'if i==0 else ('F0F5F7'if i%2 else 'FFFFFF'));tcPr.append(shd)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');mar.append(el)
   tcPr.append(mar)
   if i==0:
    for r in p.runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
  trPr=rr._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
  if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc.add_paragraph('NSS / PASS 61',style='Subtitle')
doc.add_paragraph('Can the smaller learner\nstill learn?',style='Title')
p=doc.add_paragraph('LB03-LEARNABILITY',style='Heading 1');p.paragraph_format.space_before=Pt(10)
rich(doc.add_paragraph(),'**Reconciled passes 55–59, with pass 60 as context.**\nTwo-bit forward weights; 16-bit working copies; genuine acquisition and retention measured separately.')
doc.add_paragraph('Execution v0.1  •  1 October 2026, America/Detroit\nRuntime receipts: 2 October 2026 UTC',style='Subtitle')
rich(doc.add_paragraph(),'**Decision:** Retain stochastic FP16 masters as a bounded candidate. It meets the frozen separate-regime screen, but residual failures and numerical warnings prevent broader claims.')
doc.add_paragraph('Reading guide',style='Heading 2')
for i,s in enumerate(sections,1):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);link(p,f'{i:02d}  {s["title"]}',f'sec{i:02d}')
p=doc.add_paragraph('SOURCE-BOUND • EVERY PARENT RETAINED • NO POST-LEARNING REPAIR');p.paragraph_format.space_before=Pt(14)
for r in p.runs:r.font.size=Pt(8);r.bold=True;r.font.color.rgb=RGBColor.from_string('1E6971')
for i,s in enumerate(sections,1):
 doc.add_page_break();p=doc.add_paragraph(f'{i:02d} / {s["title"]}',style='Heading 1');bookmark(p,f'sec{i:02d}')
 lines=s['body'].splitlines();ix=0
 while ix<len(lines):
  t=lines[ix].strip()
  if not t:ix+=1;continue
  if t.startswith('|'):
   arr=[]
   while ix<len(lines)and lines[ix].strip().startswith('|'):arr.append(lines[ix]);ix+=1
   table(arr);continue
  if t.startswith('#'):
   rich(doc.add_paragraph(style='Heading 2'),t.lstrip('# ').strip());ix+=1;continue
  paragraph=[t];ix+=1
  while ix<len(lines)and lines[ix].strip()and not lines[ix].strip().startswith(('|','#')):paragraph.append(lines[ix].strip());ix+=1
  rich(doc.add_paragraph(),' '.join(paragraph))
doc.core_properties.title='NSS LB03-LEARNABILITY — Source reconciliation, low-bit training, results and decision'
doc.core_properties.author='NSS Research';doc.core_properties.subject='Bounded comparison with exact source lineage and retained failures'
out=R/'docs/NSS_NBR02_LB03_LEARNABILITY_Report_v0.1.docx';doc.save(out);print(out)
